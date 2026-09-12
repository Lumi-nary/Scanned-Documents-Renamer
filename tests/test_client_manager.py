import os
import shutil
import tempfile
import unittest
from datetime import datetime
from docx import Document

from src.client_manager import (
    ClientManager,
    format_client_name,
    parse_date_to_dt,
    calculate_date_range_str,
    update_clients_docx
)

class TestClientManager(unittest.TestCase):
    def setUp(self):
        self.test_dir = tempfile.mkdtemp()
        self.watch_dir = os.path.join(self.test_dir, "Temporary", "Unproccesed")
        os.makedirs(self.watch_dir, exist_ok=True)

    def tearDown(self):
        shutil.rmtree(self.test_dir, ignore_errors=True)

    def test_format_client_name(self):
        self.assertEqual(format_client_name("BEACON HOMES DEVELOPMENT CORPORATION"), "Beacon Homes Development Corporation")
        self.assertEqual(format_client_name("MR. JOHN ROBERT SMITH"), "Mr. John Robert Smith")
        self.assertEqual(format_client_name("SUMMIT TRADE DISTRIBUTION, INC"), "Summit Trade Distribution, Inc.")
        self.assertEqual(format_client_name("ZENITH REALTY DEVELOPMENT CORP."), "Zenith Realty Development Corp")
        self.assertEqual(format_client_name(""), "General Clients")

    def test_format_client_name_acronyms_and_initials(self):
        # Sample corporations with acronyms and initials
        self.assertEqual(format_client_name("ABC MANAGEMENT CORP.,"), "ABC Management Corp")
        self.assertEqual(format_client_name("XYZ GROUP CORPORATION"), "XYZ Group Corporation")
        self.assertEqual(format_client_name("BLUESTAR J3 CORP.,"), "Bluestar J3 Corp")
        self.assertEqual(format_client_name("NEXUS WORLD MEDIA, INC."), "Nexus World Media, Inc.")
        self.assertEqual(format_client_name("CASCADE VENTURES, CORP"), "Cascade Ventures, Corp")
        self.assertEqual(format_client_name("VERTEX CONSULTANCY, OPC"), "Vertex Consultancy, OPC")
        self.assertEqual(format_client_name("ATLAS PHARMACEUTICALS, INC."), "Atlas Pharmaceuticals, Inc.")



    def test_parse_date_to_dt(self):
        dt1 = parse_date_to_dt("11/20/2018")
        self.assertIsNotNone(dt1)
        self.assertEqual(dt1.year, 2018)
        self.assertEqual(dt1.month, 11)
        self.assertEqual(dt1.day, 20)

        dt2 = parse_date_to_dt("09/20/2023")
        self.assertIsNotNone(dt2)
        self.assertEqual(dt2.year, 2023)
        self.assertEqual(dt2.month, 9)
        self.assertEqual(dt2.day, 20)

    def test_calculate_date_range_str(self):
        dates = [
            datetime(2018, 11, 20),
            datetime(2020, 5, 15),
            datetime(2023, 9, 20)
        ]
        range_str = calculate_date_range_str(dates)
        self.assertEqual(range_str, "Nov 20 2018 – Sept 20 2023")

    def test_update_clients_docx(self):
        docx_path = os.path.join(self.test_dir, "Clients.docx")
        
        # Test creation
        update_clients_docx(docx_path, "BEACON HOMES DEVELOPMENT CORPORATION", "Nov 20 2018 – Sept 20 2023")
        self.assertTrue(os.path.exists(docx_path))

        doc = Document(docx_path)
        self.assertEqual(len(doc.tables), 1)
        table = doc.tables[0]
        self.assertEqual(table.rows[0].cells[0].text, "Name")
        self.assertEqual(table.rows[0].cells[1].text, "Date Range")
        self.assertEqual(table.rows[1].cells[0].text, "Beacon Homes Development Corporation")
        self.assertEqual(table.rows[1].cells[1].text, "Nov 20 2018 – Sept 20 2023")

        # Test updating existing client row
        update_clients_docx(docx_path, "Beacon Homes Development Corporation", "Nov 20 2018 – Dec 31 2024")
        doc2 = Document(docx_path)
        table2 = doc2.tables[0]
        # Should still have 1 header row + 1 client row (updated)
        self.assertEqual(len(table2.rows), 2)
        self.assertEqual(table2.rows[1].cells[1].text, "Nov 20 2018 – Dec 31 2024")

    def test_update_clients_docx_template_empty_rows(self):
        docx_path = os.path.join(self.test_dir, "Clients.docx")
        doc = Document()
        table = doc.add_table(rows=4, cols=2)
        table.rows[0].cells[0].text = "Name"
        table.rows[0].cells[1].text = "Date Range"
        # Rows 1, 2, 3 are empty template rows
        doc.save(docx_path)

        # First update should populate row 1
        update_clients_docx(docx_path, "Mr. John Robert Smith", "Jan 10 2020 – Feb 15 2022")
        doc_updated = Document(docx_path)
        t = doc_updated.tables[0]
        self.assertEqual(len(t.rows), 4)
        self.assertEqual(t.rows[1].cells[0].text, "Mr. John Robert Smith")
        self.assertEqual(t.rows[1].cells[1].text, "Jan 10 2020 – Feb 15 2022")
        self.assertEqual(t.rows[2].cells[0].text.strip(), "")

        # Second update should populate row 2
        update_clients_docx(docx_path, "Summit Trade Distribution, Inc", "May 01 2021")
        doc_updated2 = Document(docx_path)
        t2 = doc_updated2.tables[0]
        self.assertEqual(len(t2.rows), 4)
        self.assertEqual(t2.rows[2].cells[0].text, "Summit Trade Distribution, Inc.")
        self.assertEqual(t2.rows[2].cells[1].text, "May 01 2021")

    def test_wrap_up_batch(self):
        manager = ClientManager(self.watch_dir, update_docx=True)
        
        # Create dummy file in Temporary directory
        temp_file = os.path.join(self.test_dir, "Temporary", "Retainer Services Proposal 09_22_2017.pdf")
        os.makedirs(os.path.dirname(temp_file), exist_ok=True)
        with open(temp_file, "w", encoding="utf-8") as f:
            f.write("Dummy PDF content")

        records = [
            {
                "file_path": temp_file,
                "filename": "Retainer Services Proposal 09_22_2017.pdf",
                "client_name": "BEACON HOMES DEVELOPMENT CORPORATION",
                "doc_date": "09/22/2017"
            }
        ]

        summary = manager.wrap_up_batch(records)
        self.assertEqual(len(summary), 1)
        self.assertEqual(summary[0]["client_name"], "Beacon Homes Development Corporation")
        
        expected_client_dir = os.path.join(self.test_dir, "Beacon Homes Development Corporation")
        self.assertTrue(os.path.exists(expected_client_dir))
        
        moved_file = os.path.join(expected_client_dir, "Retainer Services Proposal 09_22_2017.pdf")

        self.assertTrue(os.path.exists(moved_file))
        self.assertTrue(os.path.exists(manager.docx_path))

    def test_sticky_client_state(self):
        manager = ClientManager(self.watch_dir)
        self.assertIsNone(manager.get_active_client())

        # Doc 1 has explicit client "GLOBAL GROUP CORPORATION"
        name1, is_inherited1 = manager.resolve_client_name("GLOBAL GROUP CORPORATION")
        self.assertEqual(name1, "Global Group Corporation")
        self.assertFalse(is_inherited1)
        self.assertEqual(manager.get_active_client(), "Global Group Corporation")

        # Doc 2 has NO client (None) -> Should inherit previous client
        name2, is_inherited2 = manager.resolve_client_name(None)
        self.assertEqual(name2, "Global Group Corporation")
        self.assertTrue(is_inherited2)

        # Doc 3 has NO client ("Unrecognized") -> Should still inherit
        name3, is_inherited3 = manager.resolve_client_name("Unrecognized")
        self.assertEqual(name3, "Global Group Corporation")
        self.assertTrue(is_inherited3)

        # Doc 4 arrives with a NEW client "NEXUS MEDIA, INC." -> Changes active client
        name4, is_inherited4 = manager.resolve_client_name("NEXUS MEDIA, INC.")
        self.assertEqual(name4, "Nexus Media, Inc.")
        self.assertFalse(is_inherited4)
        self.assertEqual(manager.get_active_client(), "Nexus Media, Inc.")

        # Doc 5 has NO client -> Now inherits "Nexus Media, Inc."
        name5, is_inherited5 = manager.resolve_client_name(None)
        self.assertEqual(name5, "Nexus Media, Inc.")
        self.assertTrue(is_inherited5)

        # Explicit reset clears the active client context
        manager.reset_active_client()
        self.assertIsNone(manager.get_active_client())
        name6, is_inherited6 = manager.resolve_client_name(None)
        self.assertEqual(name6, "General Clients")
        self.assertFalse(is_inherited6)

    def test_route_file_to_client_direct(self):
        manager = ClientManager(self.watch_dir)

        # Create dummy source file in Temporary/Unproccesed
        src_file = os.path.join(self.watch_dir, "scan001.pdf")
        with open(src_file, "w", encoding="utf-8") as f:
            f.write("Dummy PDF stream")

        dest_path = manager.route_file_to_client(
            src_file=src_file,
            client_name="BLUESTAR J3 CORP.,",
            doc_date="03/28/2025",
            target_filename="Secretary's Certificate Doc. No. 74.pdf"
        )

        expected_dir = os.path.join(self.test_dir, "Bluestar J3 Corp")
        self.assertTrue(os.path.exists(expected_dir))
        self.assertTrue(os.path.exists(dest_path))
        self.assertEqual(os.path.basename(dest_path), "Secretary's Certificate Doc. No. 74.pdf")
        self.assertFalse(os.path.exists(src_file))

        # Check Clients.docx is NOT created by default
        self.assertFalse(os.path.exists(manager.docx_path))

    def test_route_file_to_client_with_docx_enabled(self):
        manager = ClientManager(self.watch_dir, update_docx=True)

        src_file = os.path.join(self.watch_dir, "scan002.pdf")
        with open(src_file, "w", encoding="utf-8") as f:
            f.write("Dummy PDF stream 2")

        dest_path = manager.route_file_to_client(
            src_file=src_file,
            client_name="BLUESTAR J3 CORP.,",
            doc_date="03/28/2025",
            target_filename="Secretary's Certificate Doc. No. 74.pdf"
        )

        self.assertTrue(os.path.exists(dest_path))
        self.assertTrue(os.path.exists(manager.docx_path))
        doc = Document(manager.docx_path)
        table = doc.tables[0]
        self.assertEqual(table.rows[1].cells[0].text, "Bluestar J3 Corp")
        self.assertEqual(table.rows[1].cells[1].text, "Mar 28 2025")

    def test_extract_client_name_corporate_secretary(self):
        import fitz
        from src.client_manager import extract_client_name_from_pdf

        # Sample 1: GLOBAL GROUP CORPORATION
        pdf1 = os.path.join(self.test_dir, "cert1.pdf")
        doc1 = fitz.open()
        page1 = doc1.new_page()
        page1.insert_textbox(fitz.Rect(50, 50, 550, 200), 'I, JUAN DELA CRUZ, Filipino, of legal age, being the Corporate Secretary of GLOBAL GROUP CORPORATION (the "Corporation"), a corporation duly organized...')
        doc1.save(pdf1)
        doc1.close()

        extracted1 = extract_client_name_from_pdf(pdf1)
        self.assertEqual(extracted1, "Global Group Corporation")

        # Sample 2: NEXUS MEDIA, INC.
        pdf2 = os.path.join(self.test_dir, "cert2.pdf")
        doc2 = fitz.open()
        page2 = doc2.new_page()
        page2.insert_textbox(fitz.Rect(50, 50, 550, 200), 'I, MARIA SANTOS, Filipino, of legal age, being the Corporate Secretary of NEXUS MEDIA, INC. (the "Corporation"), a corporation duly organized...')
        doc2.save(pdf2)
        doc2.close()

        extracted2 = extract_client_name_from_pdf(pdf2)
        self.assertEqual(extracted2, "Nexus Media, Inc.")

        # Sample 3: BLUESTAR J3 CORP.
        pdf3 = os.path.join(self.test_dir, "cert3.pdf")
        doc3 = fitz.open()
        page3 = doc3.new_page()
        page3.insert_textbox(fitz.Rect(50, 50, 550, 200), '1. I am the duly elected and qualified Corporate Secretary of BLUESTAR J3 CORP., (the "Corporation\'), a corporation duly organized...')
        doc3.save(pdf3)
        doc3.close()

        extracted3 = extract_client_name_from_pdf(pdf3)
        self.assertEqual(extracted3, "Bluestar J3 Corp")

        # Sample 4: ACME MANAGEMENT CORP. (with "of")
        pdf4 = os.path.join(self.test_dir, "cert4.pdf")
        doc4 = fitz.open()
        page4 = doc4.new_page()
        page4.insert_textbox(fitz.Rect(50, 50, 550, 200), 'I, MARIA SANTOS, Filipino citizen, of legal age, and with office address at 123 Corporate Center, Business District, Metro Manila, being the duly qualified Corporate Secretary of ACME MANAGEMENT CORP., [hereinafter "the Corporation\'], a domestic corporation duly organized...')
        doc4.save(pdf4)
        doc4.close()

        extracted4 = extract_client_name_from_pdf(pdf4)
        self.assertEqual(extracted4, "Acme Management Corp")

        # Sample 5: ACME MANAGEMENT CORP. (WITHOUT "of")
        pdf5 = os.path.join(self.test_dir, "cert5.pdf")
        doc5 = fitz.open()
        page5 = doc5.new_page()
        page5.insert_textbox(fitz.Rect(50, 50, 550, 200), 'I, MARIA SANTOS, Corporate Secretary ACME MANAGEMENT CORP., after having duly sworn to in accordance with law, do hereby depose and state: 1) That I am the duly elected, qualified and incumbent Corporate Secretary ACME MANAGEMENT CORP. (the "Corporation"), a corporation duly organized...')
        doc5.save(pdf5)
        doc5.close()

        extracted5 = extract_client_name_from_pdf(pdf5)
        self.assertEqual(extracted5, "Acme Management Corp")

        # Sample 6: Treasurer of ACME MANAGEMENT CORP. (Certification of Capital Investment)
        pdf6 = os.path.join(self.test_dir, "cert6.pdf")
        doc6 = fitz.open()
        page6 = doc6.new_page()
        page6.insert_textbox(fitz.Rect(50, 50, 550, 400), '''CERTIFICATION OF CAPITAL INVESTMENT
I, MARIA SANTOS, of legal age, with address at 123 Corporate Center, Business District, Metro Manila, after having been duly sworn in accordance with law, do hereby depose and state that:
1. I am the Treasurer of ACME MANAGEMENT CORP. (the "Corporation"), a corporation organized and existing under the laws of the Republic of the Philippines, with address at Suite 500, Innovation Tower, City Center, Philippines.
2. As Treasurer, I am authorized to receive for and on behalf of the Corporation all payments for the subscription of the Corporation.''')
        doc6.save(pdf6)
        doc6.close()

        extracted6 = extract_client_name_from_pdf(pdf6)
        self.assertEqual(extracted6, "Acme Management Corp")
        self.assertNotEqual(extracted6, "Receive For And On Behalf Of The Corporation")


if __name__ == "__main__":
    unittest.main()

